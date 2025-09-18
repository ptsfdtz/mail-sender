import openpyxl

def read_credentials(key_file):
    """读取邮箱账号和密码"""
    try:
        with open(key_file, 'r', encoding='utf-8') as file:
            lines = file.readlines()
            username = lines[0].strip()
            password = lines[1].strip()
            return username, password
    except Exception as e:
        print(f"读取密钥文件时出错：{e}")
        return None, None


def read_recipient_list(file_path):
    """读取名单 Excel，返回列表，每一项是一个 dict"""
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active

        headers = [cell.value for cell in sheet[1]]

        name_idx = headers.index("姓名")
        email_idx = headers.index("邮箱")
        dept_idx = headers.index("录取部门")
        group_idx = headers.index("群号")

        result_list = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            name = row[name_idx]
            email = row[email_idx]
            department = row[dept_idx]
            qq_group_id = row[group_idx]

            if not email:  # 跳过无邮箱的行
                continue

            result_list.append({
                "name": name,
                "email": email,
                "department": department,
                "qqGroupId": qq_group_id
            })

        return result_list

    except FileNotFoundError:
        print(f"文件未找到：{file_path}")
    except ValueError as e:
        print(f"表格缺少必要列: {e}")
    except Exception as e:
        print(f"读取文件时出错：{e}")
        return None
import openpyxl

def read_credentials(key_file):
    """读取邮箱账号和密码"""
    try:
        with open(key_file, 'r', encoding='utf-8') as file:
            lines = file.readlines()
            username = lines[0].strip()
            password = lines[1].strip()
            return username, password
    except Exception as e:
        print(f"读取密钥文件时出错：{e}")
        return None, None


def read_recipient_list(file_path):
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active

        headers = [cell.value for cell in sheet[1]]

        name_idx = headers.index("姓名")
        email_idx = headers.index("邮箱")
        dept_idx = headers.index("录取部门")
        group_idx = headers.index("群号")

        result_list = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            name = row[name_idx]
            email = row[email_idx]
            department = row[dept_idx]
            qq_group_id = row[group_idx]

            if not email:  # 跳过无邮箱的行
                continue

            result_list.append({
                "name": name,
                "email": email,
                "department": department,
                "qqGroupId": qq_group_id
            })

        return result_list

    except FileNotFoundError:
        print(f"文件未找到：{file_path}")
    except ValueError as e:
        print(f"表格缺少必要列: {e}")
    except Exception as e:
        print(f"读取文件时出错：{e}")
        return None
